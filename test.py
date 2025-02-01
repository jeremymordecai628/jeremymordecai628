from docx import Document

# Create a new document
doc = Document()

# Title
doc.add_heading('Personal Reflection on My Goals', level=1)

# Introduction
doc.add_heading('Introduction', level=2)
doc.add_paragraph("Setting goals is essential for personal and professional growth. "
                  "As a university student pursuing Information System Technology, "
                  "I recognize the importance of balancing academic, career, social, physical, "
                  "psychological, and family aspirations. This reflection outlines my short-term and "
                  "long-term goals, the steps to achieve them, and potential challenges with strategies "
                  "for overcoming them.")

# Academic Goals
doc.add_heading('1. Academic Goals', level=2)
doc.add_heading('Short-Term Goals:', level=3)
doc.add_paragraph("- Improve time management to complete assignments on time.\n"
                  "- Achieve high grades in IT-related courses.\n"
                  "- Gain proficiency in programming languages such as Python and Java.")
doc.add_heading('Long-Term Goals:', level=3)
doc.add_paragraph("- Graduate with a strong academic record.\n"
                  "- Pursue certifications in cybersecurity and cloud computing.\n"
                  "- Conduct research in emerging IT fields.")

doc.add_heading('Steps to Achieve:', level=3)
doc.add_paragraph("• Develop a structured study plan.\n"
                  "• Participate in study groups and seek mentorship.\n"
                  "• Utilize online learning platforms like Coursera and LinkedIn Learning.")

# Career Goals
doc.add_heading('2. Career Goals', level=2)
doc.add_heading('Short-Term Goals:', level=3)
doc.add_paragraph("- Secure an internship in software development or IT support.\n"
                  "- Build a strong LinkedIn profile and expand my professional network.")

doc.add_heading('Long-Term Goals:', level=3)
doc.add_paragraph("- Become a leading IT professional in cybersecurity or DevOps.\n"
                  "- Establish a tech startup focused on digital solutions.")

doc.add_heading('Steps to Achieve:', level=3)
doc.add_paragraph("• Gain hands-on experience through internships.\n"
                  "• Attend industry conferences and tech events.\n"
                  "• Work on open-source projects and contribute to GitHub.")

# Social Goals
doc.add_heading('3. Social Goals', level=2)
doc.add_paragraph("Developing strong communication skills and expanding my social network "
                  "are key social goals.")

doc.add_heading('Short-Term Goals:', level=3)
doc.add_paragraph("- Improve public speaking and presentation skills.\n"
                  "- Build meaningful friendships and professional relationships.")

doc.add_heading('Long-Term Goals:', level=3)
doc.add_paragraph("- Become an industry mentor.\n"
                  "- Expand my global professional network.")

# Physical Goals
doc.add_heading('4. Physical Goals', level=2)
doc.add_paragraph("Maintaining a healthy lifestyle is essential for sustained productivity.")

doc.add_heading('Short-Term Goals:', level=3)
doc.add_paragraph("- Exercise at least four times a week.\n"
                  "- Follow a balanced diet and stay hydrated.")

doc.add_heading('Long-Term Goals:', level=3)
doc.add_paragraph("- Maintain a consistent fitness routine.\n"
                  "- Participate in a marathon or fitness challenge.")

# Psychological Goals
doc.add_heading('5. Psychological Goals', level=2)
doc.add_heading('Short-Term Goals:', level=3)
doc.add_paragraph("- Practice mindfulness and stress management techniques.\n"
                  "- Improve emotional intelligence and self-awareness.")

doc.add_heading('Long-Term Goals:', level=3)
doc.add_paragraph("- Achieve long-term mental resilience.\n"
                  "- Support others in managing stress and mental health.")

# Family Goals
doc.add_heading('6. Family Goals', level=2)
doc.add_heading('Short-Term Goals:', level=3)
doc.add_paragraph("- Spend quality time with family despite a busy schedule.\n"
                  "- Improve communication with family members.")

doc.add_heading('Long-Term Goals:', level=3)
doc.add_paragraph("- Support family financially and emotionally.\n"
                  "- Strengthen family bonds through shared experiences.")

# Conclusion
doc.add_heading('Conclusion', level=2)
doc.add_paragraph("Setting and achieving these goals will require dedication, "
                  "perseverance, and continuous learning. Balancing academics, career, "
                  "social life, and personal well-being is crucial for long-term success. "
                  "By leveraging available resources and staying committed to growth, I am "
                  "confident in my ability to achieve these aspirations.")

# Save the document
file_path = "/mnt/data/Personal_Reflection.docx"
doc.save(file_path)

# Return the file path
file_path
